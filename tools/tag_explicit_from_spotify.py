#!/usr/bin/env python3
"""Apply Spotify's explicit flag to a FLAC library using ISRC matching.

The script is deliberately a dry run unless --write is supplied. It writes
the standard ITUNESADVISORY Vorbis comment, which Navidrome imports as its
explicitStatus field after a library rescan.

Dependencies:
    python -m pip install openpyxl mutagen

Examples:
    python tools/tag_explicit_from_spotify.py

    python tools/tag_explicit_from_spotify.py --write
"""

from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from pathlib import Path
from typing import Iterable

from mutagen.flac import FLAC
from openpyxl import load_workbook


def normalize_isrc(value: object) -> str:
    """Normalize common ISRC spellings without changing the identifier."""
    if value is None:
        return ""
    return "".join(ch for ch in str(value).strip().upper() if ch.isalnum())


def parse_explicit(value: object) -> bool | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        if value in (0, 1):
            return bool(value)
    normalized = str(value).strip().lower()
    if normalized in {"1", "true", "yes", "explicit"}:
        return True
    if normalized in {"0", "false", "no", "clean", "not explicit"}:
        return False
    raise ValueError(f"Valor de Explicit não reconhecido: {value!r}")


def column_indexes(header: Iterable[object]) -> dict[str, int]:
    return {str(value).strip(): index for index, value in enumerate(header) if value is not None}


def load_spotify_flags(path: Path) -> dict[str, bool]:
    if not path.is_file():
        raise FileNotFoundError(f"Planilha não encontrada: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        indexes = column_indexes(next(rows))
    except StopIteration:
        raise ValueError("A planilha está vazia") from None

    required = {"ISRC", "Explicit"}
    missing = required - indexes.keys()
    if missing:
        raise ValueError(f"Colunas ausentes na planilha: {', '.join(sorted(missing))}")

    flags: dict[str, bool] = {}
    duplicates: set[str] = set()
    for row_number, row in enumerate(rows, start=2):
        isrc = normalize_isrc(row[indexes["ISRC"]] if indexes["ISRC"] < len(row) else None)
        if not isrc:
            continue
        explicit = parse_explicit(row[indexes["Explicit"]] if indexes["Explicit"] < len(row) else None)
        if explicit is None:
            continue
        previous = flags.get(isrc)
        if previous is not None and previous != explicit:
            duplicates.add(isrc)
            continue
        flags[isrc] = explicit

    workbook.close()
    if duplicates:
        raise ValueError(
            "ISRCs com valores Explicit conflitantes: "
            + ", ".join(sorted(duplicates)[:10])
            + (" ..." if len(duplicates) > 10 else "")
        )
    return flags


def flac_isrcs(audio: FLAC) -> list[str]:
    values = []
    for key, raw_values in (audio.tags.items() if audio.tags else []):
        if key.lower() in {"isrc", "tsrc"}:
            values.extend(normalize_isrc(value) for value in raw_values)
    return [value for value in values if value]


def load_flac_paths(path: Path) -> list[Path]:
    """Read the exact FLAC paths from the existing library tag export."""
    if not path.is_file():
        raise FileNotFoundError(f"Exportação de tags não encontrada: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    try:
        indexes = column_indexes(next(rows))
    except StopIteration:
        raise ValueError("A exportação de tags está vazia") from None
    if "file_path" not in indexes:
        raise ValueError("A exportação de tags não possui a coluna file_path")
    paths = []
    seen: set[Path] = set()
    for row in rows:
        raw = row[indexes["file_path"]] if indexes["file_path"] < len(row) else None
        if not raw:
            continue
        file_path = Path(str(raw))
        if file_path.suffix.lower() == ".flac" and file_path not in seen:
            seen.add(file_path)
            paths.append(file_path)
    workbook.close()
    return paths


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--flac-xlsx",
        type=Path,
        default=Path(r"D:\navidrome-stream\scripts\flac_tags_export.xlsx"),
        help="Exportação com a coluna file_path (padrão: caminho usado pelo script original)",
    )
    parser.add_argument(
        "--spotify-xlsx",
        type=Path,
        default=Path(r"D:\navidrome-stream\scripts\liked_from_spotify.xlsx"),
        help="Export liked_from_spotify.xlsx (padrão: caminho usado pelo script original)",
    )
    parser.add_argument("--write", action="store_true", help="Grava ITUNESADVISORY nos FLACs; sem isso é dry-run")
    parser.add_argument("--force", action="store_true", help="Regrava tags mesmo quando o valor atual já coincide")
    args = parser.parse_args()

    try:
        spotify_flags = load_spotify_flags(args.spotify_xlsx)
        flac_paths = load_flac_paths(args.flac_xlsx)
    except (OSError, ValueError) as error:
        parser.error(str(error))

    stats = defaultdict(int)
    errors: list[str] = []
    changed: list[tuple[Path, str, bool]] = []

    for path in flac_paths:
        stats["flacs"] += 1
        if not path.is_file():
            errors.append(f"{path}: arquivo não encontrado")
            continue
        try:
            audio = FLAC(path)
            isrcs = flac_isrcs(audio)
        except Exception as error:  # mutagen can reject malformed individual files
            errors.append(f"{path}: leitura falhou: {error}")
            continue

        matches = {spotify_flags[isrc] for isrc in isrcs if isrc in spotify_flags}
        if not matches:
            stats["unmatched"] += 1
            continue
        if len(matches) > 1:
            errors.append(f"{path}: ISRCs correspondem a valores Explicit diferentes")
            continue

        explicit = matches.pop()
        desired = "1" if explicit else "0"
        current = [str(value).strip() for value in (audio.tags.get("ITUNESADVISORY", []) if audio.tags else [])]
        if current == [desired] and not args.force:
            stats["already_correct"] += 1
            continue

        stats["matched"] += 1
        changed.append((path, desired, explicit))
        if args.write:
            audio["ITUNESADVISORY"] = [desired]
            audio.save()
            stats["written"] += 1

    mode = "WRITE" if args.write else "DRY RUN"
    print(f"[{mode}] FLACs encontrados: {stats['flacs']}")
    print(f"Correspondências por ISRC: {stats['matched']}")
    print(f"Já corretos: {stats['already_correct']}")
    print(f"Sem correspondência: {stats['unmatched']}")
    if args.write:
        print(f"Gravados: {stats['written']}")
    else:
        print("Nada foi alterado. Use --write depois de revisar a prévia.")

    for path, desired, explicit in changed[:20]:
        print(f"  {'EXPLICIT' if explicit else 'CLEAN  '} -> {desired}: {path}")
    if len(changed) > 20:
        print(f"  ... e mais {len(changed) - 20} arquivo(s)")
    for error in errors[:20]:
        print(f"ERRO: {error}", file=sys.stderr)
    return 1 if errors else 0


if __name__ == "__main__":
    raise SystemExit(main())
